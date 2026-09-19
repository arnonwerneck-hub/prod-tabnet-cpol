#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Le data/tabnet_raw_long.csv (saida do extract_tabnet.py) e monta o arquivo
.xlsx final com producao das Policlinicas/Centros Cariocas por
Profissional-CBO x Mes de Atendimento, para Quantidade Apresentada e
Quantidade Aprovada, apenas 2023 em diante.
"""
import os
import re
import sys

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
IN_CSV = os.path.join(DATA_DIR, "tabnet_raw_long.csv")
OUT_XLSX = os.path.join(BASE_DIR, "Producao_Policlinicas_RJ_2023_em_diante.xlsx")

MES_ORDEM = {m: i for i, m in enumerate(["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"])}


def mes_chave(label):
    m = re.match(r"([A-Za-z]{3})/(\d{4})", label)
    if not m:
        return (9999, 99)
    mes, ano = m.group(1), int(m.group(2))
    return (ano, MES_ORDEM.get(mes, 99))


def carregar_dados():
    df = pd.read_csv(IN_CSV, encoding="utf-8")
    df = df[df["mes_atendimento"].apply(lambda x: mes_chave(x)[0] >= 2023)].copy()
    return df


def montar_pivot(df, conteudo):
    sub = df[df["conteudo"] == conteudo]
    if sub.empty:
        return pd.DataFrame()
    agg = sub.groupby(["estabelecimento_nome", "profissional_cbo", "mes_atendimento"], as_index=False)["valor"].sum()
    pivot = agg.pivot_table(
        index=["estabelecimento_nome", "profissional_cbo"],
        columns="mes_atendimento",
        values="valor",
        aggfunc="sum",
        fill_value=0,
    )
    meses_ordenados = sorted(pivot.columns, key=mes_chave)
    pivot = pivot[meses_ordenados]
    pivot["Total"] = pivot.sum(axis=1)
    pivot = pivot.reset_index()
    pivot = pivot.rename(columns={"estabelecimento_nome": "Estabelecimento", "profissional_cbo": "Profissional-CBO"})
    pivot = pivot.sort_values(["Estabelecimento", "Profissional-CBO"])
    return pivot


def montar_resumo_mensal(df):
    agg = df.groupby(["estabelecimento_nome", "mes_atendimento", "conteudo"], as_index=False)["valor"].sum()
    pivot = agg.pivot_table(index=["estabelecimento_nome", "mes_atendimento"], columns="conteudo", values="valor", fill_value=0).reset_index()
    for col in ("apresentada", "aprovada"):
        if col not in pivot.columns:
            pivot[col] = 0
    pivot["taxa_aprovacao_%"] = pivot.apply(
        lambda r: round(100.0 * r["aprovada"] / r["apresentada"], 1) if r["apresentada"] else 0.0, axis=1
    )
    pivot["ordem"] = pivot["mes_atendimento"].apply(mes_chave)
    pivot = pivot.sort_values(["estabelecimento_nome", "ordem"]).drop(columns="ordem")
    pivot = pivot.rename(
        columns={
            "estabelecimento_nome": "Estabelecimento",
            "mes_atendimento": "Mes",
            "apresentada": "Qtd_Apresentada",
            "aprovada": "Qtd_Aprovada",
        }
    )
    return pivot[["Estabelecimento", "Mes", "Qtd_Apresentada", "Qtd_Aprovada", "taxa_aprovacao_%"]]


def formatar_planilha(ws, df, freeze_cols=2):
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, col in enumerate(df.columns, start=1):
        max_len = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).values[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 10), 42)
    ws.freeze_panes = ws.cell(row=2, column=freeze_cols + 1)
    ws.auto_filter.ref = ws.dimensions


def main():
    df = carregar_dados()
    if df.empty:
        print("Nenhum dado encontrado (2023+) em", IN_CSV, file=sys.stderr)
        sys.exit(1)

    pv_apresentada = montar_pivot(df, "apresentada")
    pv_aprovada = montar_pivot(df, "aprovada")
    resumo = montar_resumo_mensal(df)

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        resumo.to_excel(writer, sheet_name="Resumo Mensal", index=False)
        pv_apresentada.to_excel(writer, sheet_name="Qtd Apresentada", index=False)
        pv_aprovada.to_excel(writer, sheet_name="Qtd Aprovada", index=False)
        df.to_excel(writer, sheet_name="Dados Brutos", index=False)

        formatar_planilha(writer.sheets["Resumo Mensal"], resumo, freeze_cols=2)
        formatar_planilha(writer.sheets["Qtd Apresentada"], pv_apresentada, freeze_cols=2)
        formatar_planilha(writer.sheets["Qtd Aprovada"], pv_aprovada, freeze_cols=2)
        formatar_planilha(writer.sheets["Dados Brutos"], df, freeze_cols=2)

    print("Arquivo salvo em:", OUT_XLSX)
    print("Linhas resumo:", len(resumo), "| Apresentada:", len(pv_apresentada), "| Aprovada:", len(pv_aprovada), "| Brutos:", len(df))


if __name__ == "__main__":
    main()
